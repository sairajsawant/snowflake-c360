"""
Build the submission deck (Snowflake 2026 brand style is primary).

    ./build_all.sh       # re-render diagrams + build -> Customer360_Decisioning_Platform_Snowflake.pptx

Content lives here; diagrams live in diagrams/*.html (edit those, never
diagrams_sf/ — it is regenerated). DECK_THEME=research python build.py still
builds the earlier research-blog style deck if ever needed.
"""
import copy, json, os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

TEMPLATE = "template.pptx"
OUT = "Customer360_Decisioning_Platform.pptx"
BG_CONTENT = "tpl/bg2.png"

# ── palette (matches diagrams/style.css) ────────────────────────────────────
INK, MUTED, RULE = "141413", "5E5D59", "E3E0D8"
IVORY, OAT = "FAF9F5", "F0EEE6"
CLAY, CLAY_T, CLAY_D = "D97757", "F5DED3", "B4583A"
SKY, SKY_T = "6A9BCC", "DDE8F3"
OLIVE, OLIVE_T = "788C5D", "E3E9D9"
KRAFT, KRAFT_T = "C49A6C", "F0E5D3"
HEATHER, HEATHER_T = "8E8BB0", "E7E6F0"

SERIF, SANS = "Georgia", "Arial"
TITLE_SIZE, TITLE_COLOR, LEDE_SIZE = 22, INK, 11.5
CARD_TITLE = INK
IMG_DIR = "img"

# ── Snowflake 2026 brand (primary). DECK_THEME=research builds the older style ─
THEME = os.environ.get("DECK_THEME", "snowflake")
if THEME == "snowflake":
    OUT = "Customer360_Decisioning_Platform_Snowflake.pptx"
    IMG_DIR = "img_sf"
    INK, MUTED, RULE = "252525", "5B5B5B", "DCE3EA"
    IVORY, OAT = "F4F8FB", "EAF2F7"
    CLAY, CLAY_T, CLAY_D = "29B5E8", "DDF2FB", "11567F"     # Snowflake Blue
    SKY, SKY_T = "11567F", "E3EDF4"                         # Mid-Blue
    OLIVE, OLIVE_T = "75CDD7", "E3F6F7"                     # Star Blue
    KRAFT, KRAFT_T = "FF9F36", "FFF0E0"                     # Valencia Orange
    HEATHER, HEATHER_T = "7254A3", "EEEAF5"                 # Purple Moon
    SERIF = SANS = "Arial"
    TITLE_SIZE, TITLE_COLOR, LEDE_SIZE = 24, "000000", 12
    CARD_TITLE = "11567F"

# official Snowflake icons, lifted as vectors from the brand template
ICON_FOR = {
    "Scattered data": "Integrated Data", "Missed warnings": "Alert", "Inconsistent action": "Consolidate",
    "Relationship Manager": "User 1", "Team Lead": "Users", "Analyst / Domain Expert": "Data Analytics",
    "Dynamic Tables": "Dynamic Tables", "Streams & Tasks": "Task", "Snowpark Python": "Snowpark",
    "Semantic View": "Metadata", "AI_COMPLETE": "LLM", "AI_SENTIMENT": "Social",
    "AI_SUMMARIZE": "Document AI", "Cortex Search": "Universal Search",
    "Cortex Agent": "Snowflake Intelligence", "Streamlit in Snowflake": "Streamlit in Snowflake",
    "Masking & row access": "Security Governance", "CoCo CLI": "Code",
    "Private by default": "Secure Data", "Scoped by role": "Role", "Spending limits": "Cost Savings",
    "Signals with evidence": "Truth", "Repeatable decisions": "Refresh", "Full audit trail": "Integrity",
    "Incremental Dynamic Tables": "Dynamic Tables", "Changed-customer streams": "Stream",
    "Daily time tick": "Time", "Signal fingerprints": "Metadata",
    "Connect real sources": "Kafka Connectors", "Actions in Slack and Jira": "Communicate",
    "Fine-grained access": "Policy", "Cross-domain, any entity": "Enterprise", "AI onboarding wizard": "Idea",
    "Predictive ML signals": "Machine Learning",
}
_ICON_SRC = None
_ICON_ID = [9000]


def icon(slide, name, x, y, size):
    """Copy a Snowflake brand icon (vector) onto the slide, fitted into a size x size box."""
    global _ICON_SRC
    if THEME != "snowflake" or name not in ICON_FOR:
        return False
    if _ICON_SRC is None:
        _ICON_SRC = (Presentation("sf_tpl.pptx"), json.load(open("sf_icons.json")))
    tpl, idx = _ICON_SRC
    sn, sid = idx[ICON_FOR[name]][:2]
    src = [sh for sh in tpl.slides[sn - 1].shapes if sh.shape_id == sid][0]
    el = copy.deepcopy(src._element)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    xfrm = el.find("p:grpSpPr/a:xfrm", ns)
    if xfrm is None:
        xfrm = el.find("p:spPr/a:xfrm", ns)
    off, ext = xfrm.find("a:off", ns), xfrm.find("a:ext", ns)
    w0, h0 = int(ext.get("cx")), int(ext.get("cy"))
    k = Inches(size) / max(w0, h0)
    w, h = int(w0 * k), int(h0 * k)
    off.set("x", str(int(Inches(x) + (Inches(size) - w) / 2)))
    off.set("y", str(int(Inches(y) + (Inches(size) - h) / 2)))
    ext.set("cx", str(w)); ext.set("cy", str(h))
    for c in el.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"):
        _ICON_ID[0] += 1
        c.set("id", str(_ICON_ID[0]))
    slide.shapes._spTree.append(el)
    return True


def rgb(h):
    return RGBColor.from_string(h)


# ── primitives ───────────────────────────────────────────────────────────────
def text(slide, x, y, w, h, paras, size=13, color=INK, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, space_after=0):
    """paras: str, or list of str / list of (text, overrides) run lists."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, o in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill=IVORY, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.9)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def dot(slide, x, y, d, fill, label=None, color="FFFFFF", size=12):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    if label:
        text(slide, x, y, d, d, label, size=size, bold=True, color=color,
             font=SERIF, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return c


def arrow(slide, x, y, w, h, fill=MUTED):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = rgb(fill); a.line.fill.background()
    return a


# ── slide frames ─────────────────────────────────────────────────────────────
def new_slide(prs, title, lede=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s.shapes.add_picture(BG_CONTENT, 0, 0, prs.slide_width, prs.slide_height)
    if THEME == "snowflake":
        box(s, 0, 0.7, 0.1, 0.36, fill=CLAY, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.45, 0.66, 9.1, 0.42, title, size=TITLE_SIZE, bold=True, font=SERIF, color=TITLE_COLOR)
    if lede:
        text(s, 0.45, 1.08, 9.1, 0.3, lede, size=LEDE_SIZE, color=MUTED)
    return s


def diagram_slide(prs, title, lede, img):
    s = new_slide(prs, title, lede)
    w = 8.75
    s.shapes.add_picture(f"{IMG_DIR}/{img}.png", Inches((10 - w) / 2), Inches(1.42), width=Inches(w))
    return s


def stat(slide, x, y, w, h, big, label, fill=IVORY, big_color=INK, big_size=28):
    if THEME == "snowflake" and big_color == INK:
        big_color = CLAY
    box(slide, x, y, w, h, fill=fill)
    text(slide, x + 0.16, y + 0.12, w - 0.3, 0.55, big, size=big_size, bold=True,
         font=SERIF, color=big_color)
    text(slide, x + 0.16, y + 0.12 + big_size / 72 * 1.25, w - 0.3, h - 0.7, label,
         size=10.5, color=MUTED)


def card(slide, x, y, w, h, title, body, fill=IVORY, accent=None, title_size=13, body_size=11):
    box(slide, x, y, w, h, fill=fill)
    tx = x + 0.18
    if icon(slide, title, x + 0.16, y + 0.1, 0.34):
        tx = x + 0.6
    elif accent:
        dot(slide, x + 0.18, y + 0.2, 0.16, accent)
        tx = x + 0.44
    text(slide, tx, y + 0.15, w - (tx - x) - 0.15, 0.3, title, size=title_size, bold=True, color=CARD_TITLE)
    text(slide, x + 0.18, y + 0.5, w - 0.33, h - 0.6, body, size=body_size, color=MUTED)


# ── build ────────────────────────────────────────────────────────────────────
prs = Presentation(TEMPLATE)

# 1 · title slide: fill the template's own fields in their own font
FIELDS = {
    "Team Name :": "Simplifiers",
    "Team Leader Name :": "Sairaj Sawant",
    "Team Size :": "1",
    "Problem Statement :": "Customer 360 and Next Best Action Engine",
}
title_slide = prs.slides[0]
for sh in title_slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() in FIELDS:
        val = FIELDS[sh.text_frame.text.strip()]
        if val:
            r0 = sh.text_frame.paragraphs[0].runs[0]
            r = sh.text_frame.paragraphs[0].add_run()
            r.text = "  " + val
            r.font.size = r0.font.size
            r.font.name = "Manrope"
            r.font.color.rgb = rgb("202729")
text(title_slide, 0.44, 3.12, 7.8, 0.32,
     [[("Customer 360 Decisioning Platform", {"bold": True, "size": 17, "font": SERIF})]])

# ---------------------------------------------------------------------------
# 2 · problem
s = new_slide(prs, "Customers signal what they'll do long before they do it",
              "Insurers and lenders already hold the evidence. It's spread across systems nobody reads together.")
pains = [
    ("Scattered data", "Policies, claims, payments, calls, emails and tickets live in separate systems.", CLAY),
    ("Missed warnings", "A customer threatens to switch insurers on a call while their claim sits stuck for weeks. Nobody links the two.", KRAFT),
    ("Inconsistent action", "Every RM decides alone, and nobody learns which actions actually work.", HEATHER),
]
for i, (t, b, c) in enumerate(pains):
    card(s, 0.45, 1.55 + i * 1.25, 5.45, 1.1, t, b, accent=c, body_size=11.5)
box(s, 6.15, 1.55, 3.4, 3.6, fill=OAT)
text(s, 6.38, 1.72, 3.0, 0.25, "WHO IT'S FOR", size=9.5, bold=True, color=MUTED)
people = [
    ("Relationship Manager", "A ranked worklist and the next best action for each customer"),
    ("Team Lead", "Approves larger offers for the team"),
    ("Analyst / Domain Expert", "Tunes signals, weights and rules without code"),
]
for i, (t, b) in enumerate(people):
    y = 2.08 + i * 0.86
    px = 6.38
    if icon(s, t, 6.36, y, 0.36):
        px = 6.84
    text(s, px, y, 9.4 - px, 0.26, t, size=13, bold=True, color=CARD_TITLE)
    text(s, px, y + 0.29, 9.4 - px, 0.5, b, size=11, color=MUTED)
text(s, 6.38, 4.66, 3.0, 0.4, "Indian health insurance and retail lending: IRDAI grievances, cashless claims, portability, EMI stress",
     size=9.5, color=MUTED, italic=True)

# ---------------------------------------------------------------------------
# 3 · solution at a glance
s = new_slide(prs, "From customer question to next best action",
              "Unify every touchpoint, read it with AI, decide with approved rules, and learn from every outcome.")
steps = [
    ("Unify", "Calls, emails, tickets, payments and policies become one live customer view.", SKY),
    ("Understand", "AI picks out intent, sentiment and evidence from every conversation.", OLIVE),
    ("Decide", "Approved rules rank the next best action for each role.", CLAY),
    ("Learn", "Every outcome updates which actions work best.", KRAFT),
]
for i, (t, b, c) in enumerate(steps):
    y = 1.58 + i * 0.92
    dot(s, 0.45, y, 0.48, c, str(i + 1), size=15)
    text(s, 1.1, y - 0.02, 3.6, 0.28, t, size=14, bold=True)
    text(s, 1.1, y + 0.28, 3.65, 0.55, b, size=11, color=MUTED)
tiles = [
    ("30", "customers unified, every one scored"), ("8", "source types, in English and Hinglish"),
    ("184", "live signals, each with its evidence"), ("3", "decision engines running side by side"),
    ("4", "reusable CoCo CLI skills"), ("19", "config rows to add a new engine"),
]
for i, (n, l) in enumerate(tiles):
    cx, cy = 5.15 + (i % 2) * 2.25, 1.55 + (i // 2) * 1.23
    stat(s, cx, cy, 2.1, 1.1, n, l, big_color=CLAY_D if i in (2, 5) else INK)

# ---------------------------------------------------------------------------
# 4-9 · architecture
diagram_slide(prs, "Architecture: four independent layers",
              "Each layer reads only from the one before it, so any layer can change without touching the rest.", "01-hld")
diagram_slide(prs, "Layer 1 · Every touchpoint, one customer",
              "A small, high-quality Indian dataset. Dynamic Tables keep it current; a semantic view gives everyone the same definitions.", "02-data-layer")
diagram_slide(prs, "Layer 2 · AI turns conversations into evidence",
              "AI extracts facts into a fixed vocabulary, and every signal keeps the sentence that justifies it.", "03-signal-layer")
diagram_slide(prs, "Layer 3 · Isolated engines on a shared vocabulary",
              "Each use case is its own engine. All read the same signals, and none can change another.", "04-decision-layer")
diagram_slide(prs, "Repeatable recommendations that keep improving",
              "AI reads, rules decide. Scoring runs on measured track records, so every answer can be reproduced and explained.", "05-learning-loop")
diagram_slide(prs, "Proactive by default, with people in the loop",
              "The platform watches every customer on a schedule and routes each decision to the person with the right authority.", "06-agentic-hitl")

# ---------------------------------------------------------------------------
# 10 · extensibility proof
s = new_slide(prs, "Adding a use case takes configuration only",
              "We added Service Recovery to the working platform. Nothing that already existed changed.")
box(s, 0.45, 1.55, 2.95, 3.6, fill=CLAY_T)
text(s, 0.68, 1.72, 2.6, 0.9, "19", size=60, bold=True, font=SERIF, color=CLAY_D)
text(s, 0.68, 2.72, 2.55, 0.3, "configuration rows", size=14, bold=True)
text(s, 0.68, 3.08, 2.55, 0.3, "0 lines of engine code", size=12, color=MUTED)
text(s, 0.68, 3.55, 2.5, 1.4, "Existing engines returned identical results before and after the change.",
     size=11, color=MUTED)
for j, (hdr, fill, rows) in enumerate([
    ("Before", OAT, ["3 customers with stuck claims were on nobody's worklist",
                     "2 customers we had let down were offered a loan top-up"]),
    ("After", OLIVE_T, ["All 3 now top the list with an expedited claim",
                        "Both get a priority callback, and no upsell until it's fixed"]),
]):
    x = 3.6 + j * 3.0
    box(s, x, 1.55, 2.85, 3.6, fill=fill)
    text(s, x + 0.2, 1.72, 2.4, 0.3, hdr.upper(), size=10, bold=True, color=MUTED)
    for k, r in enumerate(rows):
        dot(s, x + 0.2, 2.21 + k * 1.3, 0.14, CLAY if j == 0 else OLIVE)
        text(s, x + 0.46, 2.12 + k * 1.3, 2.25, 1.1, r, size=13)
arrow(s, 6.47, 3.2, 0.12, 0.28, fill=MUTED)

# ---------------------------------------------------------------------------
# 11 · CoCo CLI
diagram_slide(prs, "Built and run with CoCo CLI",
              "CoCo CLI seeded the data, built the platform, and publishes four reusable skills to Snowflake for any team.", "07-coco-skills")

# ---------------------------------------------------------------------------
# 12 · Snowflake capabilities
s = new_slide(prs, "Snowflake capabilities in use",
              "Data, AI and the app all run inside Snowflake. No customer data leaves the account.")
feats = [
    ("Dynamic Tables", "Keep the unified customer view current", SKY_T),
    ("Streams & Tasks", "Detect change and run the pipeline on schedule", SKY_T),
    ("Snowpark Python", "Portfolio-wide scoring and signal discovery", SKY_T),
    ("Semantic View", "Shared, governed business definitions", KRAFT_T),
    ("AI_COMPLETE", "Intent from calls, held to a fixed list", OLIVE_T),
    ("AI_SENTIMENT", "Tone of every conversation", OLIVE_T),
    ("AI_SUMMARIZE", "A customer's history in one paragraph", OLIVE_T),
    ("Cortex Search", "Search calls, emails and product documents", OLIVE_T),
    ("Cortex Agent", "Natural-language access to every tool", OLIVE_T),
    ("Streamlit in Snowflake", "The app runs right next to the data", KRAFT_T),
    ("Masking & row access", "PII masked, each persona sees only its book", HEATHER_T),
    ("CoCo CLI", "Built the platform and packages it as skills", HEATHER_T),
]
cw, ch = (9.1 - 3 * 0.15) / 4, 1.12
for i, (n, d, f) in enumerate(feats):
    x, y = 0.45 + (i % 4) * (cw + 0.15), 1.55 + (i // 4) * (ch + 0.15)
    box(s, x, y, cw, ch, fill=f)
    if icon(s, n, x + 0.14, y + 0.1, 0.3):
        text(s, x + 0.16, y + 0.45, cw - 0.3, 0.3, n, size=12, bold=True, color=CARD_TITLE)
        text(s, x + 0.16, y + 0.71, cw - 0.3, 0.4, d, size=10, color=MUTED)
    else:
        text(s, x + 0.16, y + 0.14, cw - 0.3, 0.3, n, size=12.5, bold=True)
        text(s, x + 0.16, y + 0.46, cw - 0.3, 0.6, d, size=10.5, color=MUTED)

# ---------------------------------------------------------------------------
# 13 · the app
s = new_slide(prs, "What's live in the app today",
              "Two modes on live Snowflake data: a guided studio and the everyday console.")
box(s, 0.45, 1.5, 4.45, 3.75, fill=OAT)
text(s, 0.68, 1.64, 4.0, 0.3, "Scenario Studio", size=14, bold=True, font=SERIF)
text(s, 0.68, 1.95, 4.0, 0.25, "The full loop on any customer, step by step", size=10.5, color=MUTED)
studio = [
    ("Stage an event", "describe it, AI writes a realistic call"),
    ("Detect", "signals extracted with their evidence"),
    ("Understand", "state recomputed, conflicts resolved"),
    ("Decide", "ranked for your role, limits checked"),
    ("Act", "carried out, customer notified, call brief"),
    ("Learn", "outcome recorded, ranking re-run"),
]
for i, (t, d) in enumerate(studio):
    y = 2.32 + i * 0.47
    dot(s, 0.68, y, 0.3, CLAY, str(i + 1), size=10)
    text(s, 1.1, y + 0.03, 3.7, 0.3, [[(t, {"bold": True}), ("  " + d, {"color": MUTED, "size": 10.5})]], size=11.5)
box(s, 5.1, 1.5, 4.45, 3.75, fill=IVORY, line=RULE)
text(s, 5.33, 1.64, 4.0, 0.3, "Operations Console", size=14, bold=True, font=SERIF)
text(s, 5.33, 1.95, 4.0, 0.25, "Each role sees only the pages it's allowed to use", size=10.5, color=MUTED)
console = [
    ("My Feed", "today's ranked worklist"), ("Decision Queue", "by severity and value"),
    ("Approvals", "one-click manager sign-off"), ("Customer 360", "next best action, act in place"),
    ("Portfolio & Learning", "value at risk, what works"), ("Ask the data", "chat that remembers context"),
    ("Config Studio", "weights, limits, rules"), ("Signal Discovery", "daily suggestions"),
]
for i, (t, d) in enumerate(console):
    y = 2.32 + i * 0.355
    dot(s, 5.33, y + 0.06, 0.11, OLIVE)
    text(s, 5.55, y, 3.9, 0.3, [[(t, {"bold": True}), ("  " + d, {"color": MUTED, "size": 10.5})]], size=11.5)

# ---------------------------------------------------------------------------
# 14 · enterprise ready
s = new_slide(prs, "Enterprise ready by design")
ent = [
    ("Private by default", "Data stays in your account, and masking policies hide PII from roles that don't need it.", SKY),
    ("Scoped by role", "Row access policies: RMs see their own book, team leads their team, analysts all.", OLIVE),
    ("Spending limits", "Offers above a role's limit wait for someone with the authority.", CLAY),
    ("Signals with evidence", "Each signal keeps the quote and confidence behind it.", KRAFT),
    ("Repeatable decisions", "Same input, same answer. AI reads, rules decide.", HEATHER),
    ("Full audit trail", "Every recommendation, action and outcome is recorded and reversible.", SKY),
]
cw = (9.1 - 2 * 0.18) / 3
for i, (t, b, c) in enumerate(ent):
    x, y = 0.45 + (i % 3) * (cw + 0.18), 1.3 + (i // 3) * 1.62
    card(s, x, y, cw, 1.47, t, b, accent=c, body_size=11.5)
box(s, 0.45, 4.6, 9.1, 0.62, fill=OAT)
text(s, 0.68, 4.6, 8.7, 0.62,
     [[("Configuration-driven.  ", {"bold": True}),
       ("Weights, limits, rules, signals and whole use cases change without a release.", {"color": MUTED})]],
     size=12, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------------------
# 15 · impact
s = new_slide(prs, "Impact", "Measured on the prototype's live data.")
imp = [
    ("93 of 184", "signals come only from what customers said, invisible to structured systems"),
    ("5", "customers caught who had fallen through the gaps"),
    ("~1 sec", "to rank every customer into today's worklist"),
    ("1 question", "from a customer question to a recommended action"),
]
for i, (n, l) in enumerate(imp):
    cx, cy = 0.45 + (i % 2) * 2.3, 1.55 + (i // 2) * 1.32
    stat(s, cx, cy, 2.15, 1.18, n, l, big_size=22, big_color=CLAY_D if i == 0 else INK)
text(s, 0.45, 4.3, 4.45, 0.9,
     [[("In production: ", {"bold": True, "color": INK}),
       ("less prep per customer, consistent actions across RMs, and a measured view of which actions work.",
        {"color": MUTED})]], size=11)
cd = CategoryChartData()
cd.categories = ["Risk", "Service", "Opportunity"]
cd.add_series("Read by AI from conversations", (82, 6, 5))
cd.add_series("Derived by rules from records", (31, 15, 45))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(5.1), Inches(1.5), Inches(4.45), Inches(3.75), cd)
ch = gf.chart
ch.has_title = True
ch.chart_title.text_frame.text = "Where the 184 signals come from"
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size = Pt(12); tp.runs[0].font.bold = True; tp.runs[0].font.name = SANS
tp.runs[0].font.color.rgb = rgb(INK)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(10); ch.legend.font.color.rgb = rgb(MUTED)
plot = ch.plots[0]
plot.gap_width = 60
plot.overlap = 100
for ser, col in zip(plot.series, (CLAY, SKY)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
    ser.format.line.color.rgb = rgb("FFFFFF")
plot.has_data_labels = True
dl = plot.data_labels
dl.font.size = Pt(10); dl.font.color.rgb = rgb("FFFFFF"); dl.font.bold = True
dl.position = XL_LABEL_POSITION.CENTER
ca, va = ch.category_axis, ch.value_axis
ca.tick_labels.font.size = Pt(11); ca.tick_labels.font.color.rgb = rgb(INK)
ca.format.line.color.rgb = rgb(RULE)
ca.reverse_order = True
va.visible = False
va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
va.major_tick_mark = XL_TICK_MARK.NONE
va.format.line.fill.background()
ca.major_tick_mark = XL_TICK_MARK.NONE
va.has_major_gridlines = False

# ---------------------------------------------------------------------------
# 16 · scalability
diagram_slide(prs, "New use cases in days, new domains in weeks",
              "Everything above the data mapping is reused. A domain expert shapes the ontology; the platform does the rest.",
              "08-onboarding-timeline")

# ---------------------------------------------------------------------------
# 16b · incremental 360
s = new_slide(prs, "An incremental 360 that scales at marginal cost",
              "Work is done only where something changed, so a bigger book doesn't mean a bigger bill.")
inc = [
    ("Incremental Dynamic Tables", "Facts refresh only where the data changed", SKY),
    ("Changed-customer streams", "Signals and state recompute only for customers who changed", OLIVE),
    ("Daily time tick", "Date-based signals touch only customers crossing a threshold today", KRAFT),
    ("Signal fingerprints", "Customers whose signals didn't change are skipped", HEATHER),
]
cw = (9.1 - 0.2) / 2
for i, (t, b, c) in enumerate(inc):
    x, y = 0.45 + (i % 2) * (cw + 0.2), 1.55 + (i // 2) * 1.32
    card(s, x, y, cw, 1.17, t, b, accent=c, title_size=14, body_size=12)
box(s, 0.45, 4.3, 9.1, 0.9, fill=CLAY_T)
text(s, 0.7, 4.3, 8.6, 0.9,
     [[("AI cost grows with new conversations.  ", {"bold": True}),
       ("Rule cost grows with changed customers, and time-based signals with threshold crossings.",
        {"color": INK})]], size=12.5, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------------------
# 17 · beyond the demo
s = new_slide(prs, "Beyond the demo", "Each step builds on the platform as it is. Estimates assume one engineer.")
nxt = [
    ("Connect real sources", "CRM, core policy and loan systems and the contact centre through Snowflake connectors, mapped once to the canonical model.", "1–2 days per source", SKY),
    ("Actions in Slack and Jira", "Approved actions post to the RM's channel and open a ticket automatically.", "2–3 days", OLIVE),
    ("Fine-grained access", "Extend row access by region and branch, and masking to every new source.", "2–3 days", HEATHER),
    ("Cross-domain, any entity", "The same engines over households, SME accounts or policies, with shared suppression across all of them.", "1–2 weeks", KRAFT),
    ("AI onboarding wizard", "Describe a use case or signal in plain English. AI drafts the configuration and a person approves it.", "1–2 weeks", CLAY),
    ("Predictive ML signals", "Snowflake ML forecasts which customers are trending toward decline and which product they'll need next, before it shows in a call.", "1–2 weeks", SKY),
]
for i, (t, d, e, c) in enumerate(nxt):
    y = 1.45 + i * 0.645
    box(s, 0.45, y, 9.1, 0.57, fill=IVORY)
    if icon(s, t, 0.58, y + 0.11, 0.34):
        text(s, 1.08, y + 0.04, 2.25, 0.5, t, size=12.5, bold=True, anchor=MSO_ANCHOR.MIDDLE, color=CARD_TITLE)
    else:
        dot(s, 0.65, y + 0.25, 0.16, c)
        text(s, 0.95, y + 0.08, 2.4, 0.5, t, size=12.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.35, y + 0.03, 4.55, 0.51, d, size=10, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 8.05, y + 0.12, 1.35, 0.33, fill=OAT, radius=0.5)
    text(s, 8.05, y + 0.12, 1.35, 0.33, e, size=9.5, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------------------
# 18 · demo
s = new_slide(prs, "What the demo shows", "One workflow end to end, input to output, in CoCo CLI and then in the app.")
cols = [
    ("INPUT", "An RM asks CoCo CLI", "“What should we do about Suresh Reddy?” and then a follow-up about another customer.", KRAFT_T),
    ("PROCESSING", "Customer Query skill", "Finds the customer, picks the one engine that answers the question, and runs it on Snowflake.", SKY_T),
    ("OUTPUT", "A grounded answer", "The ranked action with its track record and approval status, backed by his own calls.", OLIVE_T),
]
for i, (cap, t, d, f) in enumerate(cols):
    x = 0.45 + i * 3.1
    box(s, x, 1.55, 2.8, 2.35, fill=f)
    text(s, x + 0.2, 1.72, 2.4, 0.25, cap, size=9.5, bold=True, color=MUTED)
    text(s, x + 0.2, 2.02, 2.4, 0.35, t, size=14, bold=True, font=SERIF)
    text(s, x + 0.2, 2.45, 2.42, 1.4, d, size=11.5, color=INK)
    if i < 2:
        arrow(s, x + 2.86, 2.6, 0.18, 0.26)
box(s, 0.45, 4.12, 9.1, 1.08, fill=IVORY, line=RULE)
text(s, 0.68, 4.25, 8.7, 0.25, "THEN IN THE APP", size=9.5, bold=True, color=MUTED)
text(s, 0.68, 4.55, 8.7, 0.55,
     "My Feed across all three engines  ·  act straight from Customer 360  ·  a chat that remembers context  ·  "
     "the full loop in Scenario Studio  ·  one-click manager approval",
     size=12)

# ---------------------------------------------------------------------------
# reorder: title, new slides, thank-you; drop the template's guideline/blank slides
ids = prs.slides._sldIdLst
items = list(ids)
title_el, guide_els, thanks_el, new_els = items[0], items[1:5], items[5], items[6:]
for el in guide_els:
    prs.part.drop_rel(el.rId)
    ids.remove(el)
ids.remove(thanks_el)
ids.append(thanks_el)

prs.save(OUT)


def shrink_photos(path, min_bytes=1_000_000, quality=92):
    """
    The template's title and thank-you backgrounds are photographs stored as
    lossless PNG (~4.8 MB of the deck). Re-encode only those as high-quality
    JPEG — visually identical for photos — and leave the diagrams as lossless
    PNG so they stay sharp. Keeps the exported PDF well under 5 MB.
    """
    import io, re, zipfile
    from PIL import Image
    zin = zipfile.ZipFile(path)
    files = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    infos = {i.filename: i for i in zin.infolist()}
    zin.close()
    renamed = {}
    for name, data in list(files.items()):
        if name.startswith("ppt/media/") and name.endswith(".png") and len(data) > min_bytes:
            im = Image.open(io.BytesIO(data)).convert("RGB")
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=quality, subsampling=0, optimize=True)
            new = name[:-4] + ".jpeg"
            files[new] = buf.getvalue()
            del files[name]
            renamed[name.split("/")[-1]] = new.split("/")[-1]
    if not renamed:
        return
    for name in list(files):
        if name.endswith(".rels"):
            x = files[name].decode("utf8")
            for old, new in renamed.items():
                x = x.replace(f"media/{old}\"", f"media/{new}\"")
            files[name] = x.encode("utf8")
    ct = files["[Content_Types].xml"].decode("utf8")
    if 'Extension="jpeg"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="jpeg" ContentType="image/jpeg"/><Default ', 1)
    files["[Content_Types].xml"] = ct.encode("utf8")
    zout = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    order = ["[Content_Types].xml"] + [n for n in files if n != "[Content_Types].xml"]
    for n in order:
        zout.writestr(n, files[n])
    zout.close()
    print("re-encoded photos:", ", ".join(f"{k} -> {v}" for k, v in renamed.items()))


shrink_photos(OUT)
print("wrote", OUT, "with", len(prs.slides), "slides")
